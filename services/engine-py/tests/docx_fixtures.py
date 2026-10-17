"""测试夹具:python-docx 现场合成 docx(文档 RAG 测试共用)。

仅 dev 依赖(python-docx 在 dev group);只服务测试,严禁进运行时 import 链。
"""

from __future__ import annotations

import base64
import io
from pathlib import Path

# 1x1 透明 PNG(与 gateway 契约测试同源字节)
PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def build_docx(path: Path, blocks: list[tuple]) -> Path:
    """blocks 序列合成 docx。

    支持块:("h1", text) / ("h2", text) / ("h3", text) / ("p", text) /
    ("table", [[cell,...],...]) / ("image",) — image 块内嵌 1x1 PNG。
    """
    from docx import Document

    doc = Document()
    for block in blocks:
        kind, payload = (block if len(block) == 2 else (block[0], None))
        if kind in ("h1", "h2", "h3"):
            doc.add_heading(payload, level=int(kind[1]))
        elif kind == "p":
            doc.add_paragraph(payload)
        elif kind == "table":
            rows = payload
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    table.rows[i].cells[j].text = str(cell)
        elif kind == "image":
            doc.add_picture(io.BytesIO(PNG_1PX))
        else:
            raise ValueError(f"未知夹具块类型: {kind}")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
