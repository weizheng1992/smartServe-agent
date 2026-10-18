"""文档 RAG 模拟样例生成器 — docs/rag-samples/ 全部演示文档的可复现来源。

产物 = 4 份 docx(含 PIL 真画中文图表 + 表格 + 有序列表)+ 1 份内嵌图表位图的
PDF(经 Chrome headless 打印 HTML 产出)。全部内容为极光演示租户的虚构数据,
零真实业务信息(数据诚实律:演示素材同样严禁冒充真实文档)。

用法(仓库根):
    cd services/engine-py && uv run python ../../docs/rag-samples/generate_samples.py
    # PDF 步骤自动调用 Google Chrome(headless);未安装 Chrome 时按提示手动执行

再摄取入库:
    uv run --env-file .env python -m engine_py.rag.document_ingest docs/rag-samples --business-id aurora
"""

from __future__ import annotations

import base64
import subprocess
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont

OUT = Path(__file__).resolve().parent
FONT = "/System/Library/Fonts/Hiragino Sans GB.ttc"  # macOS 自带;其他平台换任意 CJK TTF
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"


def _font(size: int):
    return ImageFont.truetype(FONT, size)


def bar_chart(path: Path, title: str, labels: list[str], values: list[int], unit: str = "") -> Path:
    """中文柱状图:坐标轴 + 数值标注(VLM 读数验证的标注基准)。"""
    w, h = 860, 460
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((24, 18), title, fill="black", font=_font(26))
    left, bottom, top = 70, h - 60, 80
    max_v = max(values) * 1.15
    d.line([(left, top), (left, bottom), (w - 30, bottom)], fill="black", width=2)
    gap = (w - left - 80) / len(values)
    bw = gap * 0.55
    for i, (label, v) in enumerate(zip(labels, values)):
        x = left + gap * i + (gap - bw) / 2
        bh = (bottom - top) * v / max_v
        d.rectangle([x, bottom - bh, x + bw, bottom], fill="#4a7ebb", outline="#2f5b8f")
        d.text((x + bw / 2 - 24, bottom - bh - 26), f"{v}{unit}", fill="black", font=_font(18))
        d.text((x + bw / 2 - 18, bottom + 10), label, fill="black", font=_font(20))
    img.save(path)
    return path


def line_chart(path: Path, title: str, labels: list[str], values: list[int], unit: str = "") -> Path:
    """中文折线图:数据点 + 数值标注。"""
    w, h = 860, 460
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((24, 18), title, fill="black", font=_font(26))
    left, bottom, top = 90, h - 60, 80
    max_v = max(values) * 1.15
    d.line([(left, top), (left, bottom), (w - 30, bottom)], fill="black", width=2)
    pts = []
    for i, (label, v) in enumerate(zip(labels, values)):
        x = left + (w - left - 60) * i / (len(values) - 1)
        y = bottom - (bottom - top) * v / max_v
        pts.append((x, y))
        d.text((x - 14, bottom + 10), label, fill="black", font=_font(18))
        d.ellipse([x - 4, y - 4, x + 4, y + 4], fill="#c0392b")
        d.text((x - 34, y - 30), f"{v}{unit}", fill="#c0392b", font=_font(17))
    d.line(pts, fill="#c0392b", width=3)
    img.save(path)
    return path


def add_table(doc, rows: list[list]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            table.rows[i].cells[j].text = str(cell)


def main() -> None:
    # ── 1. 会员积分 PRD(柱状图 + 表格 + 有序列表)──
    chart1 = bar_chart(OUT / "_chart_redeem.png", "图 1  2026 年各季度积分兑换量(万分)", ["Q1", "Q2", "Q3", "Q4"], [120, 185, 240, 310])
    d1 = Document()
    d1.add_heading("极光会员积分体系产品需求文档", level=1)
    d1.add_paragraph("本 PRD 定义会员积分的获取、兑换与清零规则。背景:2025 年试点复购率提升 6.4%,2026 年全面推广。核心验收指标:兑换转化率不低于 8%,季度兑换量逐季环比正增长。")
    d1.add_heading("积分兑换门槛", level=2)
    d1.add_paragraph("积分兑换券与实物,门槛与库存策略如下:")
    add_table(d1, [["兑换项目", "所需积分", "库存策略"], ["10 元优惠券", "1000", "不限量"], ["帆布袋", "2000", "每月 500 件"], ["保温杯", "5000", "每月 200 件"], ["蓝牙音箱", "12000", "每月 60 件"]])
    d1.add_heading("季度兑换量目标", level=2)
    d1.add_paragraph("产品侧给出门店兑换量的季度目标,见下图,第四季度兑换量须达到 310 万分:")
    d1.add_picture(str(chart1), width=Inches(5.5))
    d1.add_paragraph("若连续两季未达标,须触发增长专项复盘。")
    d1.add_heading("积分清零与提醒", level=2)
    d1.add_paragraph("积分有效期截至当年 12 月 31 日,过期自动清零;清零前 30 天、7 天各推送一次提醒。等级积分与消费积分分开计算,等级积分永不过期。")
    d1.add_heading("兑换流程", level=2)
    d1.add_paragraph("用户兑换须按以下步骤闭环:")
    d1.add_paragraph("用户在积分商城选中商品", style="List Number")
    d1.add_paragraph("系统校验积分余额与库存", style="List Number")
    d1.add_paragraph("扣减积分并生成兑换码", style="List Number")
    d1.add_paragraph("实物类走物流发货,券类即时到账", style="List Number")
    d1.save(OUT / "极光会员积分体系PRD.docx")

    # ── 2. 物流配送说明(双表格 + 文字数据)──
    d2 = Document()
    d2.add_heading("极光物流配送与运费说明", level=1)
    d2.add_paragraph("本说明约束发货时效、运费计算与异常件处理,适用于极光商城全部实物订单。")
    d2.add_heading("发货时效", level=2)
    add_table(d2, [["商品类型", "发货时限", "备注"], ["现货商品", "48 小时内", "以出库扫描为准"], ["预售商品", "按页面标注", "最迟不超过 30 天"], ["定制商品", "7 个工作日", "不含法定节假日"]])
    d2.add_heading("运费标准", level=2)
    d2.add_paragraph("单笔订单满 99 元包邮;不满 99 元收取 8 元运费。偏远地区附加费如下:")
    add_table(d2, [["区域", "附加费", "预计时效"], ["新疆/西藏", "15 元", "5-7 天"], ["内蒙古/甘肃/青海", "10 元", "4-5 天"], ["其他地区", "0 元", "2-3 天"]])
    d2.add_heading("异常件处理", level=2)
    d2.add_paragraph("物流停滞超 72 小时自动升级工单;丢件按运费双倍与货值全额赔付,48 小时内到账。")
    d2.save(OUT / "极光物流配送说明.docx")

    # ── 3. 退换货政策 v2(表格 + 文字数据)──
    d3 = Document()
    d3.add_heading("极光退换货政策(2026 修订版)", level=1)
    d3.add_paragraph("v2 修订要点:退货时效从 7 天延长到 15 天,新增定制品与贴身衣物的不予退换清单。")
    d3.add_heading("退货时效", level=2)
    d3.add_paragraph("签收后 15 天内可无理由退货,吊牌完好且不影响二次销售为前提。退货入库后 3 个工作日内原路退款。")
    d3.add_heading("不予退换情形", level=2)
    add_table(d3, [["情形", "原因", "处理方式"], ["贴身衣物/袜子", "卫生考量", "不支持退换"], ["定制品", "个性定制", "不支持退换"], ["已拆封食品", "食品安全", "仅质量问题可退"], ["赠品", "非独立订单", "随主商品退回"]])
    d3.add_heading("运费承担", level=2)
    d3.add_paragraph("质量问题退货运费由平台承担;无理由退货运费由用户承担,运费 12 元封顶从退款中扣除。")
    d3.save(OUT / "极光退换货政策v2.docx")

    # ── 4. 大促作战手册(折线图 + SOP + 表格)──
    chart4 = line_chart(OUT / "_chart_gmv.png", "图 1  2025 双 11 大促 GMV 走势(百万元)", ["10.20", "10.25", "10.30", "11.01", "11.03", "11.11"], [8, 12, 18, 45, 38, 62])
    d4 = Document()
    d4.add_heading("极光大促活动作战手册", level=1)
    d4.add_paragraph("本手册约束大促期间的库存、客服与舆情三板斧,值为 2026 双 11 执行口径。")
    d4.add_heading("GMV 目标与节奏", level=2)
    d4.add_paragraph("参考 2025 年走势(见下图),11.11 当日 GMV 目标 6200 万元,预售期(10.20-10.31)蓄水占比不低于 30%:")
    d4.add_picture(str(chart4), width=Inches(5.5))
    d4.add_heading("客服 SOP", level=2)
    d4.add_paragraph("大促期间客服响应纪律:")
    d4.add_paragraph("首响时间不得超过 60 秒", style="List Number")
    d4.add_paragraph("退款类问题一律 2 小时内给出结论", style="List Number")
    d4.add_paragraph("舆情类问题 15 分钟内上报值班主管", style="List Number")
    d4.add_heading("库存红线", level=2)
    add_table(d4, [["指标", "红线", "动作"], ["可售库存", "低于 50 件", "立即下架限时购入口"], ["售罄率", "超 90%", "启动补货审批"], ["退款率", "超 8%", "暂停该品推广"]])
    d4.save(OUT / "极光大促活动作战手册.docx")

    # ── 5. 门店巡检 PDF(HTML 内嵌图表位图,Chrome headless 打印)──
    chart5 = bar_chart(OUT / "_chart_store.png", "图 1  各门店 9 月巡检达标率(%)", ["三里屯", "淮海路", "春熙路", "万象城"], [96, 88, 92, 76])
    b64 = base64.b64encode(Path(chart5).read_bytes()).decode()
    html = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><style>
body {{ font-family: "PingFang SC", "Hiragino Sans GB", sans-serif; }}
h1 {{ font-size: 22px; }} h2 {{ font-size: 16px; }}
table {{ border-collapse: collapse; }} td, th {{ border: 1px solid #333; padding: 4px 10px; }}
</style></head><body>
<h1>极光门店月度巡检报告(9 月)</h1>
<h2>一、总体结论</h2>
<p>四家直营店 9 月巡检平均达标率 88%,万象城店连续两个月低于 80% 红线,须在 10 月 15 日前完成整改复检。</p>
<h2>二、达标率明细</h2>
<img src="data:image/png;base64,{b64}" width="620"/>
<h2>三、扣分明细</h2>
<table><tr><th>门店</th><th>陈列</th><th>卫生</th><th>库存准确率</th></tr>
<tr><td>三里屯</td><td>98</td><td>95</td><td>95</td></tr>
<tr><td>淮海路</td><td>90</td><td>86</td><td>88</td></tr>
<tr><td>春熙路</td><td>93</td><td>91</td><td>92</td></tr>
<tr><td>万象城</td><td>72</td><td>80</td><td>76</td></tr></table>
<h2>四、整改要求</h2>
<p>低于 85 分的门店须提交整改计划,复核仍不达标的店长扣减当月绩效 20%。</p>
</body></html>"""
    html_path = Path("/tmp/rag_sample_report.html")
    html_path.write_text(html, encoding="utf-8")
    pdf_path = OUT / "极光门店巡检报告9月.pdf"
    if Path(CHROME).exists():
        subprocess.run(
            [CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer", f"--print-to-pdf={pdf_path}", f"file://{html_path}"],
            check=True,
            capture_output=True,
        )
    else:
        print(f"未找到 Chrome({CHROME}),请手动渲染 PDF:用任意浏览器打印 {html_path} → {pdf_path}")

    for intermediate in OUT.glob("_chart_*.png"):
        intermediate.unlink()  # 图表已内嵌进文档,中间产物不留仓库
    print("生成完成:", sorted(p.name for p in OUT.iterdir() if p.suffix in (".docx", ".pdf")))


if __name__ == "__main__":
    main()
